# -*- coding: utf-8 -*-
"""Build standalone tables/Table_S2.xlsx from the DEAS CSV (and optional PSI)."""
from __future__ import print_function

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import PSI, TABLE_S2, TABLE_S2_CSV, require

header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
header_fill = PatternFill("solid", fgColor="3C5488")
sig_fill = PatternFill("solid", fgColor="FCE4D6")
thin = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def style_header(ws):
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def load_psi(path):
    raw = path.read_text(encoding="utf-8", errors="replace").splitlines()
    hdr = raw[0].split("\t")
    if not hdr[0].startswith("ENSG") and not hdr[0].startswith("event"):
        cols = ["event_id"] + hdr
        rows = [line.split("\t") for line in raw[1:] if line.strip()]
        psi = pd.DataFrame(rows, columns=cols)
        for c in cols[1:]:
            psi[c] = pd.to_numeric(psi[c], errors="coerce")
        return psi
    return pd.read_csv(path, sep="\t")


def main():
    csv_path = require(TABLE_S2_CSV, "Table_S2_Differential_AS.csv")
    deas = pd.read_csv(csv_path)

    tumor = []
    normal = []
    psi = None
    samples = []
    if PSI.is_file():
        psi = load_psi(PSI)
        if "event_id" not in psi.columns:
            psi = psi.rename(columns={psi.columns[0]: "event_id"})
        samples = [c for c in psi.columns if c != "event_id"]
        tumor = [c for c in samples if str(c).endswith("T")]
        normal = [c for c in samples if str(c).endswith("N")]
    else:
        print("PSI not found (set PDAC_PSI_PATH); writing DEAS sheet only:", PSI)

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "README"
    notes = [
        ["Table S2. Differential alternative splicing in the JP PDAC cohort (GSE196009)"],
        [""],
        ["Sheet", "Content"],
        [
            "S2b AS Differential Events",
            "Event-level DEAS used for Figure 2b/2c. 7434 tested events; 48 pass |delta_PSI|>=0.2 and P<=0.01 (23 up, 25 down).",
        ],
        [
            "Event_PSI_JP",
            "SUPPA2 PSI matrix for the same event IDs. 13 tumor vs 6 adjacent-normal samples."
            if psi is not None
            else "Omitted: local PSI file not found.",
        ],
        [""],
        ["Source RDS", "local file via PDAC_RDS_PATH (not in this repo)"],
        ["Source PSI", str(PSI) if psi is not None else "not available"],
        ["Threshold in paper", "|delta_PSI| >= 0.2 and unadjusted P <= 0.01 (not FDR)."],
        ["Tumor samples (n=13)", ", ".join(tumor) if tumor else "see PSI file"],
        ["Normal samples (n=6)", ", ".join(normal) if normal else "see PSI file"],
        [""],
        [
            "Does NOT replace Figure 2a counts",
            "Paper Figure 2a cites 10,871 novel + 12,643 known events. This file is the DEAS subset, not that landscape catalog.",
        ],
    ]
    for r in notes:
        ws0.append(r)
    ws0["A1"].font = Font(bold=True, name="Arial", size=14, color="3C5488")
    ws0.column_dimensions["A"].width = 28
    ws0.column_dimensions["B"].width = 120
    for cell in ws0[3]:
        cell.font = header_font
        cell.fill = header_fill

    ws1 = wb.create_sheet("S2b AS Differential Events")
    for r_i, row in enumerate(dataframe_to_rows(deas, index=False, header=True), 1):
        ws1.append(list(row))
        if r_i == 1:
            continue
            is_sig = bool(deas.iloc[r_i - 2]["Meets paper threshold"])
        for c_i, cell in enumerate(ws1[r_i], 1):
            cell.font = Font(name="Arial", size=9)
            cell.border = thin
            if is_sig:
                cell.fill = sig_fill
            hdr = deas.columns[c_i - 1]
            if hdr in ("Mean PSI (tumor)", "Mean PSI (normal)", "Delta PSI", "mean_PSI_tumor", "mean_PSI_normal", "delta_PSI") and isinstance(cell.value, float):
                cell.number_format = "0.000"
            if hdr in ("P value", "P value (BH-adjusted)", "P_value", "P_adj_BH") and isinstance(cell.value, float):
                cell.number_format = "0.00E+00"
    style_header(ws1)
    ws1.column_dimensions["A"].width = 72
    for col in range(2, deas.shape[1] + 1):
        ws1.column_dimensions[get_column_letter(col)].width = 14
    ws1.column_dimensions["C"].width = 14
    ws1.column_dimensions["D"].width = 16

    if psi is not None:
        ws2 = wb.create_sheet("Event_PSI_JP")
        group = ["group"] + ["tumor" if c in tumor else ("normal" if c in normal else "") for c in samples]
        ws2.append(["event_id"] + samples)
        for r_i, row in enumerate(dataframe_to_rows(psi[["event_id"] + samples], index=False, header=False), 2):
            ws2.append(list(row))
            for cell in ws2[r_i]:
                cell.font = Font(name="Arial", size=8)
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
        style_header(ws2)
        ws2.column_dimensions["A"].width = 72
        ws2.insert_rows(2)
        ws2["A2"] = "sample_group"
        for i, g in enumerate(group[1:], 2):
            cell = ws2.cell(2, i, g)
            cell.font = Font(bold=True, name="Arial", size=9, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="E64B35" if g == "tumor" else "4DBBD5")
            cell.alignment = Alignment(horizontal="center")
        ws2.freeze_panes = "B3"
        ws2.auto_filter.ref = "A1:{}{}".format(get_column_letter(1 + len(samples)), ws2.max_row)

    wb.save(TABLE_S2)
    print("saved", TABLE_S2, TABLE_S2.stat().st_size)
    print("DEAS", deas.shape, "sig", int(deas["Meets paper threshold"].sum()))
    if psi is not None:
        print("PSI", psi.shape, "T", len(tumor), "N", len(normal), tumor, normal)


if __name__ == "__main__":
    main()
