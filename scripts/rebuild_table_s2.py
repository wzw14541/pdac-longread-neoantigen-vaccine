# -*- coding: utf-8 -*-
"""Replace Table S2 with previously computed JP-cohort DEAS results."""
from __future__ import print_function

import shutil
import sys
from pathlib import Path

import pandas as pd
import pyreadr
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import (
    RDS,
    SUPP_ORIGINAL,
    SUPP_TABLES,
    TABLE_S2_CSV,
    TABLE_S2_NOTE,
    TABLES,
)

COLMAP = [
    ("evt", "event_id"),
    ("gene_id", "gene_id"),
    ("gn_name", "gene_symbol"),
    ("ok_name", "event_label"),
    ("as_tp", "AS_type"),
    ("nvl", "novel_isoform_event"),
    ("t_psi", "mean_PSI_tumor"),
    ("n_psi", "mean_PSI_normal"),
    ("dpsi", "delta_PSI"),
    ("raw_p", "P_value"),
    ("p_adj", "P_adj_BH"),
    ("label", "significance"),
]


def main():
    if not RDS.is_file():
        raise SystemExit(
            "Missing RDS: {}\nSet PDAC_RDS_PATH to jp_evt_rst.rds".format(RDS)
        )

    xlsx = SUPP_TABLES if SUPP_TABLES.is_file() else SUPP_ORIGINAL
    if not xlsx.is_file():
        raise SystemExit("Missing supplementary workbook: {}".format(xlsx))
    backup = TABLES / "Supplementary_Tables_RealPDAC_revised_backup_before_S2.xlsx"

    print("Reading", RDS)
    df = pyreadr.read_r(str(RDS))[None].reset_index(drop=True)
    out = pd.DataFrame({new: df[old] for old, new in COLMAP})
    out["paper_threshold"] = (out["P_value"] <= 0.01) & (out["delta_PSI"].abs() >= 0.2)
    out["direction"] = "NS"
    out.loc[out["paper_threshold"] & (out["delta_PSI"] > 0), "direction"] = "Up"
    out.loc[out["paper_threshold"] & (out["delta_PSI"] < 0), "direction"] = "Down"
    out["abs_delta_PSI"] = out["delta_PSI"].abs()
    out = out.sort_values(
        ["paper_threshold", "abs_delta_PSI", "P_value"],
        ascending=[False, False, True],
    ).drop(columns=["abs_delta_PSI"])
    n_sig = int(out["paper_threshold"].sum())
    n_up = int((out["direction"] == "Up").sum())
    n_down = int((out["direction"] == "Down").sum())
    print("events", len(out), "paper_threshold", n_sig, "up", n_up, "down", n_down)

    published = out.rename(columns={
        "event_id": "Event ID",
        "gene_id": "Gene ID",
        "gene_symbol": "Gene symbol",
        "event_label": "Event label",
        "AS_type": "AS type",
        "novel_isoform_event": "Novel-isoform event",
        "mean_PSI_tumor": "Mean PSI (tumor)",
        "mean_PSI_normal": "Mean PSI (normal)",
        "delta_PSI": "Delta PSI",
        "P_value": "P value",
        "P_adj_BH": "P value (BH-adjusted)",
        "significance": "Significance",
        "paper_threshold": "Meets paper threshold",
        "direction": "Direction",
    })
    published.to_csv(TABLE_S2_CSV, index=False)
    print("wrote", TABLE_S2_CSV)

    TABLE_S2_NOTE.write_text(
        "Table S2b is the JP / GSE196009 DEAS result (13 tumor vs 6 normal).\n"
        "Source RDS (local, not in this repo): set PDAC_RDS_PATH.\n"
        "PSI matrix (local, not in this repo): set PDAC_PSI_PATH.\n"
        "Paper threshold: |delta_PSI| >= 0.2 and P <= 0.01 -> 48 events (23 up, 25 down).\n"
        "Official workbook: tables/Supplementary_Tables.xlsx (sheets S2a / S2b / S2c).\n"
        "CSV: tables/Table_S2_Differential_AS.csv\n"
        "Standalone S2b+S2c: tables/Table_S2.xlsx\n",
        encoding="utf-8",
    )

    if not backup.exists() and SUPP_ORIGINAL.is_file():
        shutil.copy2(SUPP_ORIGINAL, backup)
        print("backup", backup)
    else:
        print("backup already exists or original missing")

    print("Opening workbook...")
    wb = load_workbook(xlsx)
    print("sheets", wb.sheetnames)

    old_name = "Table_S2_PDAC_PSI"
    keep_name = "S2a Isoform PSI matrix"
    new_name = "S2b AS Differential Events"
    if old_name in wb.sheetnames and keep_name not in wb.sheetnames:
        wb[old_name].title = keep_name
        print("renamed", old_name, "->", keep_name)
    if "Table_S2_Differential_AS" in wb.sheetnames:
        del wb["Table_S2_Differential_AS"]
        print("removed previous Table_S2_Differential_AS")
    if new_name in wb.sheetnames:
        del wb[new_name]
        print("removed previous", new_name)

    ws = wb.create_sheet(new_name, 1)
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    header_fill = PatternFill("solid", fgColor="3C5488")
    sig_fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for r_i, row in enumerate(dataframe_to_rows(published, index=False, header=True), 1):
        ws.append(row)
        if r_i == 1:
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
        else:
            is_sig = bool(published.iloc[r_i - 2]["Meets paper threshold"])
            for c_i, cell in enumerate(ws[r_i], 1):
                cell.font = Font(name="Arial", size=9)
                cell.border = thin
                if is_sig:
                    cell.fill = sig_fill
                hdr = published.columns[c_i - 1]
                if hdr in ("Mean PSI (tumor)", "Mean PSI (normal)", "Delta PSI") and isinstance(cell.value, float):
                    cell.number_format = "0.000"
                if hdr in ("P value", "P value (BH-adjusted)") and isinstance(cell.value, float):
                    cell.number_format = "0.00E+00"

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    widths = {
        "A": 78, "B": 22, "C": 14, "D": 16, "E": 10,
        "F": 20, "G": 16, "H": 16, "I": 12, "J": 12,
        "K": 12, "L": 16, "M": 16, "N": 12,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 22

    out_xlsx = SUPP_TABLES
    wb.save(out_xlsx)
    print("saved", out_xlsx, "sheets", wb.sheetnames)
    print("DONE")


if __name__ == "__main__":
    main()
